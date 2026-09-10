import json
import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook
from agent.research_export import export_mapping
from agent.research_quality import audit_rows
from agent.settings import load_config


class ExportTests(unittest.TestCase):
    def test_two_sheets_match_template_and_keep_sources(self):
        columns=load_config()['mapping_columns']
        row={'公司名称':'Example','业务描述':'Packaging','信源链接':'https://example.com/round','信源时间':'2026-09-01','字段证据':json.dumps({'业务描述':['S1']})}
        evidence={'S1':{'link':'https://example.com/round','title':'Example raises funding','publish_date':'2026-09-01'}}
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'Mapping.xlsx'
            export_mapping(path, columns, [row], [row], evidence)
            wb=load_workbook(path)
            self.assertEqual(wb.sheetnames,['确定性高','待核验'])
            for sheet in wb:
                self.assertEqual([c.value for c in sheet[1]],columns)
                self.assertEqual(sheet.max_column,12)
                self.assertEqual(sheet.max_row, 3)
                self.assertTrue(sheet['A3'].value.startswith('备注：'))
                self.assertEqual(sheet.auto_filter.ref, 'A1:L2')
                self.assertFalse(any(cell.comment for row in sheet for cell in row))
            wb.close()
            from scripts.merge_mapping_files import read_workbook
            from agent.io.excel import read_sheet_as_dicts
            _, records = read_workbook(path)
            self.assertEqual(len(records), 2)
            self.assertEqual(len(read_sheet_as_dicts(path)), 1)

    def test_old_chinese_date_is_not_current_proof(self):
        row={'公司名称':'Example','融资时间':'2020年','字段证据':json.dumps({'上市状态':['S1']})}
        accepted,review=audit_rows([row],[],{'S1':{'title':'Example funding','link':'https://example.com','content':'Seed funding'}})
        self.assertFalse(accepted)
        self.assertIn('超过两年',review[0]['核验状态'])

    def test_generic_article_cannot_establish_high_confidence(self):
        row={'公司名称':'Invented Corp','字段证据':json.dumps({'上市状态':['S1']})}
        accepted,review=audit_rows([row],[],{'S1':{'title':'Global market outlook','link':'https://example.com/report','content':'Invented Corp'}})
        self.assertFalse(accepted)
        self.assertIn('独立报道',review[0]['核验状态'])


if __name__=='__main__':
    unittest.main()
