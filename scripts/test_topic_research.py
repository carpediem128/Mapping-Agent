"""Focused regression tests for research provenance and Excel safety."""
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import load_workbook
from agent.topic_research import validate_company, job_path, write_book, start


class ResearchTests(unittest.TestCase):
    def test_unsupported_fields_stay_blank(self):
        evidence = {'S1': {'link': 'https://example.com/news', 'content': 'Example develops chips'}}
        row = validate_company({'fields': {'公司名称': 'Example', '公司团队': 'Invented founder'}, 'citations': {'公司名称': ['S1'], '公司团队': ['FAKE']}}, evidence)
        self.assertEqual(row['公司名称'], 'Example')
        self.assertNotIn('公司团队', row)
        self.assertEqual(row['信源链接'], 'https://example.com/news')

    def test_no_company_without_source(self):
        self.assertIsNone(validate_company({'fields': {'公司名称': 'Made up'}, 'citations': {}}, {}))

    def test_path_traversal(self):
        with self.assertRaises(ValueError):
            job_path('../config')

    def test_empty_topic(self):
        with self.assertRaises(ValueError):
            start(' ')

    def test_no_key(self):
        with patch('agent.topic_research.GLMClient.from_config', return_value=None):
            with self.assertRaises(ValueError):
                start('TGV')

    def test_excel_literal_text_and_schema(self):
        with tempfile.TemporaryDirectory() as d:
            path = Path(d) / 'test.xlsx'
            write_book(path, [('一级市场', ['公司名称','业务描述'], [{'公司名称': '=HYPERLINK("bad")', '业务描述': 'chips'}])])
            wb = load_workbook(path)
            ws = wb.active
            self.assertEqual(ws.max_column, 2)
            self.assertNotEqual(ws['A2'].data_type, 'f')
            self.assertEqual(ws.freeze_panes, 'C2')
            wb.close()


if __name__ == '__main__':
    unittest.main()
