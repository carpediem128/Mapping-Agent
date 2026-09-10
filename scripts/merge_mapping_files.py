from __future__ import annotations

import json
import re
import shutil
import sqlite3
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from agent.json_utils import dumps_json


ROOT = Path(__file__).resolve().parents[1]
WEB_DIR = ROOT / "web"
DATA_DIR = WEB_DIR / "data"
SOURCE_PATH = DATA_DIR / "mapping_sources.json"
SITE_PATH = DATA_DIR / "site_data.json"
DB_PATH = DATA_DIR / "mapping.db"


def clean(value: Any) -> str:
    return str(value if value is not None else "").replace("\r\n", "\n").strip()


def source_label(path: Path) -> str:
    stem = path.stem
    stem = re.sub(r"_?Mapping$", "", stem, flags=re.I)
    return stem.replace("_", " ").strip() + " Mapping"


def company_value(fields: dict[str, str]) -> str:
    for key, value in fields.items():
        if "公司名称" in key or "企业名称" in key:
            return value
    return fields.get("公司", "")


def read_workbook(path: Path) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    label = source_label(path)
    records: list[dict[str, Any]] = []
    sources: list[dict[str, Any]] = []
    workbook = load_workbook(path, read_only=True, data_only=True)
    for worksheet in workbook.worksheets:
        rows = worksheet.iter_rows(values_only=True)
        first = next(rows, None)
        if not first:
            continue
        headers = [clean(value) or f"未命名列{index + 1}" for index, value in enumerate(first)]
        source = {"source_id": str(path), "source_label": label, "source_file": path.name, "sheet": worksheet.title, "headers": headers, "row_count": 0}
        sources.append(source)
        for row_number, values in enumerate(rows, start=2):
            fields = {header: clean(values[index] if index < len(values) else "") for index, header in enumerate(headers)}
            if not any(fields.values()):
                continue
            topic = fields.get(headers[0], "")
            company = company_value(fields)
            if not company and topic.startswith('备注：'):
                continue
            records.append({
                "id": f"{label}|{worksheet.title}|{row_number}",
                "source_label": label,
                "source_file": path.name,
                "sheet": worksheet.title,
                "row_number": row_number,
                "topic": topic,
                "company": company,
                "fields": fields,
                "search_text": " ".join([label, worksheet.title, topic, company, *fields.values()]).lower(),
                "quality_flags": [],
            })
            source["row_count"] += 1
    workbook.close()
    return sources, records


def _ensure_db(db: sqlite3.Connection) -> None:
    db.executescript("""
    CREATE TABLE IF NOT EXISTS mapping_workbooks (
      id INTEGER PRIMARY KEY, source_key TEXT NOT NULL UNIQUE, source_file TEXT NOT NULL,
      source_label TEXT NOT NULL, workflow_status TEXT NOT NULL, uploaded_at TEXT NOT NULL,
      updated_at TEXT NOT NULL, sheet_count INTEGER NOT NULL DEFAULT 0, record_count INTEGER NOT NULL DEFAULT 0
    );
    CREATE TABLE IF NOT EXISTS mapping_records (
      id TEXT PRIMARY KEY, workbook_id INTEGER NOT NULL, source_file TEXT NOT NULL,
      sheet_name TEXT NOT NULL, market_level TEXT NOT NULL, row_number INTEGER NOT NULL,
      topic TEXT, company_name TEXT, fields_json TEXT NOT NULL, search_text TEXT NOT NULL,
      quality_flags_json TEXT NOT NULL DEFAULT '[]'
    );
    CREATE INDEX IF NOT EXISTS idx_mapping_company ON mapping_records(company_name);
    CREATE INDEX IF NOT EXISTS idx_mapping_topic ON mapping_records(topic);
    CREATE INDEX IF NOT EXISTS idx_mapping_level ON mapping_records(market_level);
    """)


def _upsert_db(sources: list[dict[str, Any]], records: list[dict[str, Any]]) -> dict[str, Any]:
    now = datetime.now().isoformat(timespec="seconds")
    db = sqlite3.connect(DB_PATH)
    db.execute("PRAGMA journal_mode=WAL")
    _ensure_db(db)
    source_keys: dict[str, int] = {}
    for item in sources:
        key = f"{item['source_file']}|{item['source_label']}"
        db.execute("""
        INSERT INTO mapping_workbooks(source_key, source_file, source_label, workflow_status, uploaded_at, updated_at)
        VALUES (?, ?, ?, 'mapping_completed', ?, ?)
        ON CONFLICT(source_key) DO UPDATE SET source_file=excluded.source_file, source_label=excluded.source_label,
          workflow_status='mapping_completed', updated_at=excluded.updated_at
        """, (key, item["source_file"], item["source_label"], now, now))
        source_keys[key] = db.execute("SELECT id FROM mapping_workbooks WHERE source_key = ?", (key,)).fetchone()[0]
    for record in records:
        key = f"{record['source_file']}|{record['source_label']}"
        workbook_id = source_keys[key]
        db.execute("""
        INSERT INTO mapping_records(id, workbook_id, source_file, sheet_name, market_level, row_number, topic, company_name, fields_json, search_text, quality_flags_json)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(id) DO UPDATE SET fields_json=excluded.fields_json, search_text=excluded.search_text,
          quality_flags_json=excluded.quality_flags_json, workbook_id=excluded.workbook_id
        """, (record["id"], workbook_id, record["source_file"], record["sheet"],
              "secondary" if record["sheet"] == "二级市场" else "primary", record["row_number"], record["topic"],
              record["company"], json.dumps(record["fields"], ensure_ascii=False), record["search_text"], json.dumps(record["quality_flags"], ensure_ascii=False)))
    db.execute("UPDATE mapping_workbooks SET sheet_count=(SELECT COUNT(DISTINCT sheet_name) FROM mapping_records WHERE workbook_id=mapping_workbooks.id), record_count=(SELECT COUNT(*) FROM mapping_records WHERE workbook_id=mapping_workbooks.id)")
    db.commit()
    result = {"workbooks": db.execute("SELECT COUNT(*) FROM mapping_workbooks").fetchone()[0], "records": db.execute("SELECT COUNT(*) FROM mapping_records").fetchone()[0], "primary_records": db.execute("SELECT COUNT(*) FROM mapping_records WHERE market_level='primary'").fetchone()[0], "secondary_records": db.execute("SELECT COUNT(*) FROM mapping_records WHERE market_level='secondary'").fetchone()[0]}
    db.close()
    return result


def merge_mapping_files(paths: list[str | Path], run_id: str | None = None) -> dict[str, Any]:
    input_paths = [Path(path).resolve() for path in paths]
    all_sources: list[dict[str, Any]] = []
    all_records: list[dict[str, Any]] = []
    for path in input_paths:
        sources, records = read_workbook(path)
        all_sources.extend(sources)
        all_records.extend(records)

    existing = json.loads(SOURCE_PATH.read_text(encoding="utf-8")) if SOURCE_PATH.exists() else {"sources": [], "records": []}
    source_files = {path.name for path in input_paths}
    sources = [item for item in existing.get("sources", []) if item.get("source_file") not in source_files] + all_sources
    records = [item for item in existing.get("records", []) if item.get("source_file") not in source_files] + all_records
    source_payload = {"generated_at": datetime.now().isoformat(timespec="seconds"), "source_count": len({item["source_file"] for item in sources}), "sheet_count": len(sources), "record_count": len(records), "sources": sources, "records": records}
    SOURCE_PATH.write_text(dumps_json(source_payload, ensure_ascii=False, indent=2), encoding="utf-8")

    artifact_dir = WEB_DIR / "artifacts" / (run_id or f"mapping_import_{datetime.now().strftime('%Y%m%d_%H%M%S')}") / "mappings"
    artifact_dir.mkdir(parents=True, exist_ok=True)
    site = json.loads(SITE_PATH.read_text(encoding="utf-8")) if SITE_PATH.exists() else {}
    mappings = [item for item in site.get("mappings", []) if item.get("name", "").split(" · ")[0] not in {item["source_label"] for item in all_sources}]
    for path in input_paths:
        label = source_label(path)
        shutil.copy2(path, artifact_dir / path.name)
        for source in [item for item in all_sources if item["source_label"] == label]:
            rows = [item for item in all_records if item["source_label"] == label and item["sheet"] == source["sheet"]]
            contact_count = sum(1 for item in rows if any("联系方式" in key and value for key, value in item["fields"].items()))
            mappings.append({"sector": label, "row_count": len(rows), "contact_count": contact_count, "preview": [], "topics": [], "sheets": [source], "secondary_count": len(rows) if source["sheet"] == "二级市场" else 0, "secondary_preview": [], "name": f"{label} · {source['sheet']}", "url": str((artifact_dir / path.name).relative_to(WEB_DIR)).replace("\\", "/")})
    site["mappings"] = mappings
    site["generated_at"] = datetime.now().isoformat(timespec="seconds")
    site.setdefault("stats", {})["mapping_count"] = len(mappings)
    site["stats"]["mapping_rows"] = sum(int(item.get("row_count") or 0) for item in mappings)
    site["reference_workbooks"] = {"source_count": source_payload["source_count"], "sheet_count": source_payload["sheet_count"], "record_count": source_payload["record_count"], "sources": sources}
    serialized = dumps_json(site, ensure_ascii=False, indent=2)
    SITE_PATH.write_text(serialized, encoding="utf-8")
    (DATA_DIR / "site_data.js").write_text("window.__MAPPING_SITE_DATA__ = " + serialized + "\n", encoding="utf-8")
    return {"source_files": [path.name for path in input_paths], "sheets": len(all_sources), "records_added": len(all_records), "database": _upsert_db(all_sources, all_records), "run_id": run_id or ""}


if __name__ == "__main__":
    import sys

    print(json.dumps(merge_mapping_files(sys.argv[1:]), ensure_ascii=False, indent=2))
