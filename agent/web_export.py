from __future__ import annotations

import json
import shutil
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from agent.json_utils import dumps_json
from agent.settings import ROOT
from agent.sources.market import DEMO_SUBTOPICS


def export_run_to_web(run_dir: Path, web_dir: Path | None = None) -> dict[str, Any]:
    """Export one agent run into static assets consumed by the web dashboard."""
    run_dir = run_dir.resolve()
    web_dir = (web_dir or ROOT / "web").resolve()
    data_dir = web_dir / "data"
    artifact_root = web_dir / "artifacts" / run_dir.name
    data_dir.mkdir(parents=True, exist_ok=True)
    artifact_root.mkdir(parents=True, exist_ok=True)

    manifest = _read_json(run_dir / "run_manifest.json")
    reports = _copy_files(run_dir / "reports", artifact_root / "reports")
    hotspots = _load_hotspots(run_dir, manifest)
    topic_by_sector = {
        str(item.get("name") or ""): item.get("subtopics") or []
        for item in hotspots.get("top20", [])
    }
    for path in (run_dir / "mappings").glob("*.xlsx"):
        sector = _mapping_sector(path)
        topic_by_sector.setdefault(sector, DEMO_SUBTOPICS.get(sector, []))
    mappings = _summarize_mappings(run_dir / "mappings", artifact_root / "mappings", topic_by_sector)
    contacts = _summarize_contact_book(run_dir, manifest)
    industry_stats = _build_industry_stats(hotspots, mappings)
    topic_count = sum(len(item.get("topics", [])) for item in mappings)

    payload = {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "run_id": run_dir.name,
        "manifest": manifest,
        "hotspots": hotspots,
        "industry_stats": industry_stats,
        "contacts": contacts,
        "reports": reports,
        "mappings": mappings,
        "stats": {
            "report_count": len(reports),
            "hotspot_count": len(hotspots.get("top20", [])),
            "mapping_count": len(mappings),
            "mapping_rows": sum(item["row_count"] for item in mappings),
            "topic_count": topic_count,
            "contact_rows": contacts.get("rows", 0),
            "contact_with_value": contacts.get("with_contact", 0),
            "contact_rate": contacts.get("contact_rate", 0),
            "industry_count": len(industry_stats),
        },
    }

    serialized = dumps_json(payload, ensure_ascii=False, indent=2)
    (data_dir / "site_data.json").write_text(serialized, encoding="utf-8")
    (data_dir / "site_data.js").write_text(
        "window.__MAPPING_SITE_DATA__ = " + serialized + "\n",
        encoding="utf-8",
    )
    return payload


def _read_json(path: Path) -> dict[str, Any]:
    if not path.exists():
        return {}
    return json.loads(path.read_text(encoding="utf-8-sig"))


def _copy_files(src_dir: Path, dst_dir: Path) -> list[dict[str, str]]:
    if not src_dir.exists():
        return []
    dst_dir.mkdir(parents=True, exist_ok=True)
    copied: list[dict[str, str]] = []
    for src in sorted(src_dir.glob("*")):
        if not src.is_file():
            continue
        dst = dst_dir / src.name
        shutil.copy2(src, dst)
        copied.append({"name": src.name, "type": src.suffix.lower().lstrip(".") or "file", "url": _web_path(dst)})
    return copied


def _load_hotspots(run_dir: Path, manifest: dict[str, Any]) -> dict[str, Any]:
    configured = str(manifest.get("hotspot", {}).get("structured_output") or "").strip()
    candidates = [Path(configured)] if configured else []
    candidates.extend(sorted(run_dir.glob("reports/*二级市场热点列表.json")))
    for path in candidates:
        if path.exists():
            return _read_json(path)
    return {
        "date": run_dir.name[:8],
        "market_summary": "本次运行没有结构化热点产出。",
        "data_status": "missing",
        "top20": [],
        "top10": [],
    }


def _summarize_mappings(src_dir: Path, dst_dir: Path, topic_by_sector: dict[str, list[dict[str, Any]]]) -> list[dict[str, Any]]:
    if not src_dir.exists():
        return []
    dst_dir.mkdir(parents=True, exist_ok=True)
    items: list[dict[str, Any]] = []
    for src in sorted(src_dir.glob("*.xlsx")):
        dst = dst_dir / src.name
        shutil.copy2(src, dst)
        sector = _mapping_sector(src)
        summary = _inspect_mapping(src, topic_by_sector.get(sector, []), sector)
        summary.update({"name": src.name, "url": _web_path(dst)})
        items.append(summary)
    return items


def _inspect_mapping(path: Path, topic_defs: list[dict[str, Any]], sector: str | None = None) -> dict[str, Any]:
    """Inspect both primary-market and secondary-market sheets by header schema."""
    wb = load_workbook(path, read_only=True, data_only=True)
    sheet_summaries: list[dict[str, Any]] = []
    primary_ws = None
    secondary_ws = None
    for ws in wb.worksheets:
        headers = [str(cell.value or "").strip() for cell in ws[1]]
        rows = _sheet_records(ws, headers)
        sheet_summaries.append({"name": ws.title, "headers": headers, "row_count": len(rows)})
        header_set = set(headers)
        if "公司名称" in header_set and ("联系方式" in header_set or "融资轮次" in header_set):
            primary_ws = primary_ws or (ws, headers, rows)
        if "公司名称" in header_set and ("股票代码" in header_set or any("市值" in header for header in headers)):
            secondary_ws = secondary_ws or (ws, headers, rows)
    primary_ws = primary_ws or (wb.active, [str(cell.value or "").strip() for cell in wb.active[1]], _sheet_records(wb.active, [str(cell.value or "").strip() for cell in wb.active[1]]))
    _, primary_headers, primary_rows = primary_ws
    records = [_primary_record(row) for row in primary_rows]
    secondary_records = [_secondary_record(row) for row in (secondary_ws[2] if secondary_ws else [])]
    wb.close()

    contact_count = sum(1 for row in records if row["contact"])
    topic_views = []
    for topic in topic_defs:
        aliases = [str(value).strip() for value in topic.get("aliases", []) if str(value).strip()]
        matched = [row for row in records if any(alias.casefold() in _record_blob(row).casefold() for alias in aliases)]
        topic_views.append({
            "name": str(topic.get("name") or ""),
            "aliases": aliases,
            "summary": str(topic.get("summary") or ""),
            "row_count": len(matched),
            "preview": matched[:6],
        })

    return {
        "sector": sector or _mapping_sector(path),
        "row_count": len(records),
        "contact_count": contact_count,
        "preview": records[:5],
        "topics": topic_views,
        "sheets": sheet_summaries,
        "secondary_count": len(secondary_records),
        "secondary_preview": secondary_records[:5],
    }


def _sheet_records(ws: Any, headers: list[str]) -> list[dict[str, str]]:
    records: list[dict[str, str]] = []
    for row_idx in range(2, ws.max_row + 1):
        row = {headers[col - 1]: _cell(ws, row_idx, col) for col in range(1, ws.max_column + 1) if col <= len(headers) and headers[col - 1]}
        if any(row.values()):
            records.append(row)
    return records


def _primary_record(row: dict[str, str]) -> dict[str, str]:
    return {
        "company": row.get("公司名称", ""),
        "region": row.get("地区", ""),
        "founded_at": row.get("成立时间", ""),
        "business": row.get("业务描述", ""),
        "product": row.get("核心产品/技术", ""),
        "round": row.get("融资轮次", ""),
        "amount": row.get("最新融资金额", ""),
        "investors": row.get("主要投资方", ""),
        "team": row.get("公司团队", ""),
        "description": row.get("公司简介", ""),
        "contact": row.get("联系方式", ""),
    }


def _secondary_record(row: dict[str, str]) -> dict[str, str]:
    return {
        "company": row.get("公司名称", ""),
        "ticker": row.get("股票代码", ""),
        "exchange": row.get("交易所", ""),
        "region": row.get("地区", ""),
        "business": row.get("业务描述", ""),
        "product": row.get("核心产品/技术", ""),
        "business_share": row.get("波浪能业务占比", ""),
        "market_cap": next((value for key, value in row.items() if "市值" in key), ""),
        "management": row.get("创始人/管理层", ""),
        "customers": row.get("主要客户", ""),
        "description": row.get("公司简介", ""),
    }


def _record_blob(row: dict[str, str]) -> str:
    return " ".join(row.values())
def _summarize_contact_book(run_dir: Path, manifest: dict[str, Any]) -> dict[str, Any]:
    configured = str(manifest.get("bd_followup", {}).get("output") or "").strip()
    candidates = [Path(configured)] if configured else []
    candidates.extend(sorted(run_dir.glob("reports/*硬科技项目总表-BD跟进.xlsx")))
    for path in candidates:
        if path.exists():
            return _inspect_contact_book(path)
    return {
        "source": "",
        "rows": 0,
        "with_contact": 0,
        "contact_rate": 0,
        "contacted": 0,
        "meetings": 0,
        "sector_breakdown": [],
        "preview": [],
    }


def _inspect_contact_book(path: Path) -> dict[str, Any]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    headers = [str(cell.value or "").strip() for cell in ws[1]]
    rows: list[dict[str, str]] = []
    for row_idx in range(2, ws.max_row + 1):
        row = {
            headers[col - 1]: str(ws.cell(row_idx, col).value or "").strip()
            for col in range(1, ws.max_column + 1)
            if headers[col - 1]
        }
        if any(row.values()):
            rows.append(row)
    wb.close()

    with_contact = sum(1 for row in rows if row.get("联系方式"))
    contacted = sum(1 for row in rows if row.get("是否联系上") == "是")
    meetings = sum(1 for row in rows if row.get("是否已约见") == "是")
    grouped: dict[str, dict[str, int]] = {}
    for row in rows:
        sector = row.get("所属赛道") or "未分类"
        item = grouped.setdefault(sector, {"rows": 0, "with_contact": 0})
        item["rows"] += 1
        if row.get("联系方式"):
            item["with_contact"] += 1
    sector_breakdown = [
        {
            "sector": sector,
            "rows": values["rows"],
            "with_contact": values["with_contact"],
            "rate": round(values["with_contact"] / values["rows"] * 100) if values["rows"] else 0,
        }
        for sector, values in sorted(grouped.items(), key=lambda pair: (-pair[1]["rows"], pair[0]))
    ]
    preview = [
        {
            "sector": row.get("所属赛道", ""),
            "company": row.get("公司名称", ""),
            "contact": row.get("联系方式", ""),
            "reached": row.get("是否联系上", "否"),
            "meeting": row.get("是否已约见", "否"),
        }
        for row in rows[:10]
    ]
    return {
        "source": path.name,
        "rows": len(rows),
        "with_contact": with_contact,
        "contact_rate": round(with_contact / len(rows) * 100) if rows else 0,
        "contacted": contacted,
        "meetings": meetings,
        "sector_breakdown": sector_breakdown,
        "preview": preview,
    }


def _build_industry_stats(hotspots: dict[str, Any], mappings: list[dict[str, Any]]) -> list[dict[str, Any]]:
    industry_by_sector = {
        str(item.get("name") or ""): str(item.get("primary_industry") or "硬科技")
        for item in hotspots.get("top20", [])
    }
    grouped: dict[str, dict[str, Any]] = {}
    for item in hotspots.get("top20", []):
        industry = str(item.get("primary_industry") or "硬科技")
        group = grouped.setdefault(industry, {"industry": industry, "hotspots": 0, "score_total": 0, "mapping_rows": 0, "contact_count": 0})
        group["hotspots"] += 1
        group["score_total"] += int(item.get("score") or 0)
    for item in mappings:
        industry = industry_by_sector.get(str(item.get("sector") or ""), "硬科技")
        group = grouped.setdefault(industry, {"industry": industry, "hotspots": 0, "score_total": 0, "mapping_rows": 0, "contact_count": 0})
        group["mapping_rows"] += int(item.get("row_count") or 0)
        group["contact_count"] += int(item.get("contact_count") or 0)
    result = []
    for group in grouped.values():
        group["avg_score"] = round(group["score_total"] / group["hotspots"]) if group["hotspots"] else 0
        group["contact_rate"] = round(group["contact_count"] / group["mapping_rows"] * 100) if group["mapping_rows"] else 0
        del group["score_total"]
        result.append(group)
    return sorted(result, key=lambda item: (-item["mapping_rows"], -item["avg_score"], item["industry"]))


def _column_index(headers: list[str], header: str) -> int | None:
    return headers.index(header) + 1 if header in headers else None


def _cell(ws: Any, row_idx: int, col_idx: int | None) -> str:
    if not col_idx:
        return ""
    return str(ws.cell(row_idx, col_idx).value or "").strip()


def _cell_by_header(ws: Any, row_idx: int, headers: list[str], header: str) -> str:
    return _cell(ws, row_idx, _column_index(headers, header))


def _mapping_sector(path: Path) -> str:
    wb = load_workbook(path, read_only=True, data_only=True)
    title = str(wb.active.title or "").strip()
    wb.close()
    if title and title.lower() not in {"sheet", "sheet1"}:
        return title
    return _sector_from_mapping_name(path.stem)

def _sector_from_mapping_name(stem: str) -> str:
    parts = stem.split("_")
    return parts[1] if len(parts) >= 3 else stem


def _web_path(path: Path) -> str:
    web_root = ROOT / "web"
    return path.resolve().relative_to(web_root.resolve()).as_posix()








