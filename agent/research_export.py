"""Two-sheet export with the original 12-column mapping schema."""
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


def export_mapping(path, columns, primary, review, evidence):
    from agent.topic_research import write_book
    groups = [('确定性高', columns, primary), ('待核验', columns, review)]
    write_book(path, groups)
    wb = load_workbook(path)
    for title, _, rows in groups:
        ws = wb[title]
        footer = len(rows) + 2
        ws.merge_cells(start_row=footer, start_column=1, end_row=footer, end_column=len(columns))
        cell = ws.cell(footer, 1)
        cell.value = '备注：本表依据公开信息生成，未知字段留空，确定性高不等同于已完成尽调，待核验项目不计入覆盖统计，来源与核验依据保留在研究任务记录中。'
        cell.font = Font(name='Microsoft YaHei', size=11, color='52616B')
        cell.fill = PatternFill('solid', fgColor='F0F5F5')
        cell.alignment = Alignment(vertical='center', wrap_text=True)
        ws.row_dimensions[footer].height = 32
    wb.save(path)
    wb.close()
