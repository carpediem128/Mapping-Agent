from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def write_rows_xlsx(path: Path, sheet_name: str, headers: list[str], rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]
    ws.append(headers)
    for row in rows:
        ws.append([row.get(header, "") for header in headers])
    _format_sheet(ws, len(headers))
    wb.save(path)


def read_sheet_as_dicts(path: Path, sheet_name: str | None = None) -> list[dict[str, Any]]:
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet_name] if sheet_name else _select_data_sheet(wb)
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        wb.close()
        return []
    headers = [str(x or "").strip() for x in rows[0]]
    out: list[dict[str, Any]] = []
    for row in rows[1:]:
        item = {headers[i]: row[i] if i < len(row) else "" for i in range(len(headers)) if headers[i]}
        if any(str(v or "").strip() for v in item.values()):
            out.append(item)
    wb.close()
    return out


def _select_data_sheet(wb: Any) -> Any:
    """Choose the primary company sheet instead of trusting workbook order."""
    preferred = None
    fallback = wb.active
    for ws in wb.worksheets:
        headers = {str(cell.value or "").strip() for cell in ws[1]}
        if "公司名称" in headers and "联系方式" in headers:
            preferred = ws
            break
        if "公司名称" in headers and preferred is None:
            preferred = ws
    return preferred or fallback


def update_contact_column(path: Path, contacts: dict[str, str]) -> int:
    wb = load_workbook(path)
    ws = _select_data_sheet(wb)
    headers = [str(ws.cell(1, c).value or "").strip() for c in range(1, ws.max_column + 1)]
    if "公司名称" not in headers or "联系方式" not in headers:
        wb.close()
        raise ValueError(f"{path} 缺少 公司名称 或 联系方式 列")
    company_col = headers.index("公司名称") + 1
    contact_col = headers.index("联系方式") + 1
    updated = 0
    for row_idx in range(2, ws.max_row + 1):
        company = str(ws.cell(row_idx, company_col).value or "").strip()
        if company in contacts:
            ws.cell(row_idx, contact_col).value = contacts[company]
            updated += 1
    _format_sheet(ws, ws.max_column)
    wb.save(path)
    wb.close()
    return updated


def _format_sheet(ws: Any, header_count: int) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    for col in range(1, header_count + 1):
        cell = ws.cell(1, col)
        cell.fill = fill
        cell.font = font
        width = max(12, min(36, len(str(cell.value or "")) * 2 + 4))
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"
